from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document

from .schema import Scene, Segment, Shot, Storyboard


def read_input(path: str | Path) -> tuple[str, Storyboard | None]:
    input_path = Path(path)
    suffix = input_path.suffix.lower()
    if suffix == ".json":
        storyboard = Storyboard.model_validate_json(input_path.read_text(encoding="utf-8"))
        return storyboard.source or storyboard.title, storyboard
    if suffix == ".docx":
        storyboard = _try_read_storyboard_docx(input_path)
        if storyboard:
            return storyboard.source or _storyboard_text(storyboard), storyboard
        return _read_docx_text(input_path), None
    if suffix in {".txt", ".md"}:
        return input_path.read_text(encoding="utf-8"), None
    raise ValueError(f"Unsupported input type: {input_path.suffix}")


def _read_docx_text(path: Path) -> str:
    doc = Document(path)
    chunks: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            chunks.append(text)
    for table in doc.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                chunks.append(" | ".join(values))
    return "\n".join(chunks)


def _try_read_storyboard_docx(path: Path) -> Storyboard | None:
    doc = Document(path)
    title = "Educational Video"
    scenes: list[Scene] = []
    current_scene: Scene | None = None
    pending_narration = False
    shot_refs: list[tuple[int, int, str]] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if text.lower().startswith("storyboard:"):
            title = text.replace("Storyboard:", "", 1).strip() or title
            continue
        if text.lower().endswith(" storyboard") and not scenes:
            title = text[:-len(" storyboard")].strip() or title
            continue
        scene_match = re.match(r"scene\s+(\d+)\s*:\s*(.+)", text, re.IGNORECASE)
        if scene_match:
            current_scene = Scene(
                scene_number=int(scene_match.group(1)),
                title=scene_match.group(2).strip(),
                narration="",
                segments=[],
            )
            scenes.append(current_scene)
            pending_narration = False
            continue
        if text.lower().startswith("narration/audio/voiceover"):
            pending_narration = True
            continue
        shot_match = re.match(r"shot\s+(\d+)\.(\d+)\s*:\s*(.+)", text, re.IGNORECASE)
        if shot_match:
            shot_refs.append((int(shot_match.group(1)), int(shot_match.group(2)), shot_match.group(3).strip()))
            continue
        if pending_narration and current_scene is not None:
            current_scene.narration = text.strip().strip('"')
            pending_narration = False

    if not scenes:
        return None

    vertical_tables = [table for table in doc.tables if _is_vertical_shot_table(table)]
    if vertical_tables:
        scenes_by_number = {scene.scene_number: scene for scene in scenes}
        for table, (scene_number, segment_number, heading) in zip(vertical_tables, shot_refs):
            scene = scenes_by_number.get(scene_number)
            if scene is not None:
                scene.segments.append(_segment_from_vertical_table(table, segment_number, heading))
        for scene in scenes:
            if not scene.segments and scene.narration:
                scene.segments = _segments_from_text(scene.narration)
        return Storyboard(title=title, source=_storyboard_text_from_scenes(scenes), scenes=scenes)

    scene_index = 0
    for table in doc.tables:
        while scene_index < len(scenes) and scenes[scene_index].segments:
            scene_index += 1
        if scene_index >= len(scenes):
            break
        scene = scenes[scene_index]
        headers = [_normalize_header(cell.text) for cell in table.rows[0].cells]
        def col(*names: str) -> int | None:
            wanted = {_normalize_header(name) for name in names}
            for index, header in enumerate(headers):
                if header in wanted:
                    return index
            return None
        number_i = col("S.no", "S no", "No") or 0
        narration_i = col("Splitting the Narration (Sentence wise)", "Narration", "Sentence") or 1
        template_i = col("Template", "Shot Template", "Visual Type", "visual_type")
        heading_i = col("Heading", "Title", "Shot Title")
        tool_i = col("Tool", "Render Tool")
        media_i = col("Media Type")
        motion_i = col("Motion Plan", "Motion / Subtitle", "Motion")
        subtitle_i = col("Subtitle", "Caption", "Subtitle Text")
        visual_i = col("Visual / Animation", "Visual", "Animation", "Visual Subject")
        local_i = col("Local Teaching Animation")
        labels_i = col("Labels", "Educational Labels")
        placement_i = col("Label Placement", "label_placement", "Placement")
        style_i = col("Label Style", "label_style")
        overlay_i = col("Overlay Plan", "Labels / Overlay Plan", "Label Plan")
        image_i = col("Image Recommendation", "Asset / Image Prompt", "Image Prompt", "Image Requirement", "image_requirement")
        prompt_i = col("ComfyUI Prompt", "ComfyUI / Image Prompt")
        formula_i = col("Formula", "Formula Lines", "Derivation")
        steps_i = col("Steps", "Explain Steps")
        wan_i = col("Wan Video Shot", "LTX Video Shot")
        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) <= max(number_i, narration_i) or not cells[number_i].strip().isdigit():
                continue
            sentence = cells[narration_i].strip().strip('"')
            template = _cell(cells, template_i)
            heading = _cell(cells, heading_i)
            tool = _cell(cells, tool_i)
            media_type = _cell(cells, media_i)
            motion = _cell(cells, motion_i)
            subtitle = _cell(cells, subtitle_i)
            visual_main = _cell(cells, visual_i)
            local = _cell(cells, local_i)
            overlay_plan = _cell(cells, overlay_i)
            label_placement = _cell(cells, placement_i)
            label_style = _cell(cells, style_i)
            labels = _labels(_cell(cells, labels_i) or overlay_plan)
            image_recommendation = _cell(cells, image_i)
            comfy_prompt = _cell(cells, prompt_i)
            formula_text = _cell(cells, formula_i)
            steps_text = _cell(cells, steps_i)
            wan_prompt = _cell(cells, wan_i)
            visual = visual_main or local or image_recommendation or sentence
            planning_text = " | ".join(part for part in [media_type, motion, visual_main, local] if part)
            image_text = wan_prompt or comfy_prompt or image_recommendation or visual_main or local or visual or sentence
            keywords = labels or _keywords(sentence + " " + visual + " " + image_text)
            shot = _shot_from_storyboard(
                template=template,
                heading=heading,
                tool=tool,
                media_type=media_type,
                motion=motion,
                visual=visual,
                image_prompt=image_text,
                labels=labels,
                overlay_plan=overlay_plan,
                label_placement=label_placement,
                label_style=label_style,
                formula_text=formula_text,
                steps_text=steps_text,
                narration=subtitle or sentence,
            )
            scene.segments.append(
                Segment(
                    segment_number=int(cells[number_i].strip()),
                    narration=subtitle or sentence,
                    visual=visual,
                    image_prompt=image_text,
                    keywords=keywords,
                    shot=shot,
                )
            )
        scene_index += 1

    for scene in scenes:
        if not scene.segments and scene.narration:
            scene.segments = _segments_from_text(scene.narration)

    return Storyboard(title=title, source=_storyboard_text_from_scenes(scenes), scenes=scenes)


def _is_vertical_shot_table(table) -> bool:
    if not table.rows or any(len(row.cells) != 2 for row in table.rows):
        return False
    keys = {_normalize_header(row.cells[0].text) for row in table.rows}
    return "narration" in keys and bool(keys & {
        "template media", "visual type image requirement", "shot id", "visual type"
    })


def _segment_from_vertical_table(table, segment_number: int, shot_heading: str) -> Segment:
    fields = {
        _normalize_header(row.cells[0].text): row.cells[1].text.strip()
        for row in table.rows if len(row.cells) >= 2
    }
    sentence = fields.get("narration", "").strip().strip('"')
    if "shot id" in fields or "visual type" in fields:
        return _segment_from_production_fields(fields, segment_number, shot_heading)
    template_lines = _plain_items(fields.get("template media", ""))
    template = _normalize_template_value(template_lines[0] if template_lines else "")
    media_type = template_lines[1] if len(template_lines) > 1 else ""
    tool_lines = _plain_items(fields.get("tool motion type", ""))
    tool = tool_lines[0] if tool_lines else ""

    visual_sections = _named_sections(fields.get("visual type image requirement", ""))
    visual_type = visual_sections.get("visual type", "")
    visual = visual_sections.get("image requirement", "") or visual_sections.get("visual notes", "") or sentence
    local_animation = visual_sections.get("local animation", "")

    asset_sections = _named_sections(fields.get("asset image prompt", ""))
    image_prompt = asset_sections.get("image prompt", "") or visual
    motion_sections = _named_sections(fields.get("motion subtitle", ""))
    motion = motion_sections.get("motion", "") or local_animation
    subtitle = motion_sections.get("subtitle", "")
    subtitle_style = motion_sections.get("subtitle style", "")
    overlay_plan = "\n".join(
        value for value in [motion_sections.get("arrows", ""), motion_sections.get("highlights", "")]
        if value
    )

    labels_text = fields.get("labels", "")
    labels = _labels(labels_text)
    label_placement = fields.get("label placement", "")
    label_style = fields.get("label style", "")

    quality_sections = _named_sections(fields.get("coverage asset quality", ""))
    formula_text = quality_sections.get("formula lines", "")
    steps_text = quality_sections.get("explain steps", "") or quality_sections.get("steps", "")

    wan_sections = _named_sections(fields.get("wan video shot", ""))
    ltx_sections = _named_sections(fields.get("ltx video shot", ""))
    video_prompt = ltx_sections.get("prompt", "") or wan_sections.get("prompt", "")
    if video_prompt:
        image_prompt = video_prompt

    # The vertical format carries both the legacy renderer template and the
    # authoritative visual type. Prefer the visual type so an unlabeled
    # realistic image cannot be routed as a labeled image merely because an
    # older Template / Media value says "labeled_image".
    routing_template = _normalize_template_value(visual_type) or template
    shot = _shot_from_storyboard(
        template=routing_template,
        heading=shot_heading,
        tool=tool,
        media_type=media_type,
        motion=motion,
        visual=visual,
        image_prompt=image_prompt,
        labels=labels,
        overlay_plan=overlay_plan,
        label_placement=label_placement,
        label_style=label_style,
        formula_text=formula_text,
        steps_text=steps_text,
        narration=sentence,
    )
    if shot is not None:
        if visual_type:
            shot.motion["visual_type"] = visual_type
        if subtitle:
            shot.motion["subtitle"] = subtitle
        if subtitle_style:
            shot.motion["subtitle_style"] = subtitle_style

    return Segment(
        segment_number=segment_number,
        narration=sentence,
        visual=visual,
        image_prompt=image_prompt,
        keywords=labels or _keywords(sentence + " " + visual + " " + image_prompt),
        shot=shot,
    )


def _segment_from_production_fields(
    fields: dict[str, str], segment_number: int, shot_heading: str
) -> Segment:
    sentence = fields.get("narration", "").strip().strip('"')
    visual_type = fields.get("visual type", "").strip()
    media_type = fields.get("media type", "").strip()
    visual = fields.get("image requirement", "").strip() or sentence
    image_prompt = fields.get("image prompt", "").strip() or visual
    labels = _labels(fields.get("labels", ""))
    label_placement = fields.get("label placement", "")
    label_style = fields.get("label style", "")
    motion = fields.get("motion", "")
    subtitle = fields.get("subtitle", "")
    subtitle_style = fields.get("subtitle style", "")
    asset_path = fields.get("asset path", "").strip() or None
    wan_prompt = fields.get("wan video prompt", "").strip()
    ltx_prompt = fields.get("ltx video prompt", "").strip()
    review_notes = fields.get("review notes", "").strip()

    routing_template = _normalize_template_value(visual_type)
    video_prompt = ltx_prompt or wan_prompt
    if video_prompt and routing_template == "short_motion_clip":
        image_prompt = video_prompt
    shot = _shot_from_storyboard(
        template=routing_template,
        heading=shot_heading,
        tool="",
        media_type=media_type,
        motion=motion,
        visual=visual,
        image_prompt=image_prompt,
        labels=labels,
        overlay_plan="",
        label_placement=label_placement,
        label_style=label_style,
        formula_text="",
        steps_text="",
        narration=sentence,
    )
    if shot is not None:
        shot.asset_path = asset_path
        shot.motion["visual_type"] = visual_type
        shot.motion["subtitle"] = subtitle
        shot.motion["subtitle_style"] = subtitle_style
        shot.motion["review_notes"] = review_notes
        if wan_prompt:
            shot.motion["wan_video_prompt"] = wan_prompt
        if ltx_prompt:
            shot.motion["ltx_video_prompt"] = ltx_prompt

    return Segment(
        segment_number=segment_number,
        narration=sentence,
        visual=visual,
        image_prompt=image_prompt,
        keywords=labels or _keywords(sentence + " " + visual + " " + image_prompt),
        shot=shot,
    )


def _normalize_template_value(text: str) -> str:
    value = _normalize_header(text).replace(" ", "_")
    aliases = {
        "video_b_roll": "video_broll",
        "split_screen_comparison": "split_screen",
        "realistic_image": "photo",
        "diagram_overlay": "diagram_overlay",
        "process_steps": "process_steps",
        "realistic_labeled_image": "realistic_labeled_image",
        "realistic_background_with_labels": "realistic_background_with_labels",
        "short_motion_clip": "short_motion_clip",
        "photo_with_labels": "labeled_image",
        "animation_with_labels": "labeled_image",
    }
    return aliases.get(value, value)


def _named_sections(text: str) -> dict[str, str]:
    known = (
        "visual_type", "image_requirement", "visual notes", "local animation",
        "asset_path", "image recommendations", "image_prompt", "motion", "subtitle",
        "subtitle_style", "arrows", "highlights", "asset quality", "formula lines",
        "explain steps", "steps", "columns", "template", "heading", "duration",
        "clip duration", "clip count", "join", "prompt", "negative",
    )
    pattern = re.compile(
        r"(?im)^(" + "|".join(re.escape(name) for name in known) + r")\s*:\s*"
    )
    matches = list(pattern.finditer(text or ""))
    result: dict[str, str] = {}
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        key = _normalize_header(match.group(1))
        result[key] = text[match.end():end].strip()
    return result


def _normalize_header(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def _cell(cells: list[str], index: int | None) -> str:
    if index is None or index >= len(cells):
        return ""
    return _clean_storyboard_cell(cells[index])


def _clean_storyboard_cell(text: str) -> str:
    text = (text or "").strip()
    text = re.sub(
        r"(?im)^\s*(visual[_ ]subject|asset[_ /-]*image[_ ]prompt|image[_ ]prompt|"
        r"overlay[_ ]plan|motion[_ /-]*subtitle|template|tool|heading)\s*:\s*",
        "",
        text,
    )
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _labels(text: str) -> list[str]:
    text = _extract_labeled_section(text or "")
    labels: list[str] = []
    for part in re.split(r"[\n,;]+", text or ""):
        clean = _clean_storyboard_cell(part).strip()
        clean = re.sub(r"(?i)^(labels?|arrows?|legend|callouts?)\s*:\s*", "", clean).strip()
        if clean and not _is_empty_overlay_noise(clean):
            labels.append(clean)
    return labels


def _extract_labeled_section(text: str) -> str:
    match = re.search(
        r"(?is)\blabels?\s*:\s*(.*?)(?=\b(?:arrows?|highlights?|formula\s+lines?|explain\s+steps?|steps?|columns?)\s*:|$)",
        text or "",
    )
    return match.group(1).strip() if match else text


def _is_empty_overlay_noise(text: str) -> bool:
    normalized = re.sub(r"[^a-z]+", " ", (text or "").lower()).strip()
    return normalized in {
        "",
        "arrows",
        "highlights",
        "formula lines",
        "explain steps",
        "steps",
        "columns",
        "explain",
        "pointing from",
        "pollinator explain",
        "style indicating",
    }


def _two_steps(text: str) -> list[str]:
    clauses = [part.strip(" .") for part in re.split(r",|\band\b|\bthen\b|\bwhile\b", text or "") if part.strip()]
    if len(clauses) >= 2:
        return [_short_title(clauses[0], "Observe"), _short_title(clauses[1], "Connect")]
    return [_short_title(text, "Observe"), "Connect to the concept"]


def _shot_from_storyboard(
    *,
    template: str,
    heading: str,
    tool: str,
    media_type: str,
    motion: str,
    visual: str,
    image_prompt: str,
    labels: list[str],
    overlay_plan: str,
    label_placement: str,
    label_style: str,
    formula_text: str,
    steps_text: str,
    narration: str,
) -> Shot | None:
    routing_text = " ".join([template, tool, media_type, motion, visual, image_prompt, overlay_plan, formula_text]).lower()
    clean_template = _normalize_header(template).replace(" ", "_")
    display_subject = _display_subject(visual or image_prompt or narration)
    shot_heading = _short_title(heading or display_subject or narration, "Key Concept")
    motion_data = _motion_dict(motion)
    if label_style:
        motion_data["label_style"] = label_style

    if clean_template in {"title", "title_card", "intro", "intro_card"} or "title card" in routing_text:
        return Shot(
            template="title_card",
            heading=shot_heading,
            subheading=_short_title(display_subject or narration, ""),
            motion=motion_data,
        )

    if clean_template in {"photo", "realistic_image"}:
        return Shot(template="photo", heading=shot_heading, motion=motion_data)

    has_formula_text = bool(_formula_lines(formula_text))
    formula_requested = clean_template in {"formula", "formula_card", "derivation"} or any(
        term in " ".join([template, media_type, tool]).lower()
        for term in ["formula", "derivation", "equation"]
    )
    if formula_requested and (has_formula_text or formula_text or clean_template in {"formula", "formula_card", "derivation"}):
        formula_lines = _formula_lines(formula_text or visual)
        explain_steps = _plain_items(steps_text) or _explain_steps(narration)
        return Shot(
            template="formula",
            heading=shot_heading,
            formula_lines=formula_lines or [_short_title(visual or narration, "Formula")],
            explain_steps=explain_steps,
            motion=motion_data,
        )

    if clean_template in {"split_screen", "split", "comparison"} or "split screen" in routing_text:
        return Shot(
            template="split_screen",
            heading=shot_heading,
            columns=_columns_from_text(overlay_plan or visual or image_prompt),
            motion=motion_data,
        )

    if clean_template in {"video_broll", "ltx", "ltx_video", "wan", "wan_video", "short_motion_clip"} or "ltx" in routing_text or "wan video" in routing_text:
        return Shot(template="short_motion_clip" if clean_template == "short_motion_clip" else "video_broll", heading=shot_heading, motion=motion_data)

    label_templates = {
        "labeled_image",
        "labelled_image",
        "labels",
        "diagram_labels",
        "diagram_overlay",
        "realistic_labeled_image",
        "realistic_background_with_labels",
        "process_steps",
    }
    if clean_template in label_templates or labels or "label" in routing_text:
        template_name = clean_template if clean_template in {
            "diagram_overlay", "realistic_labeled_image", "realistic_background_with_labels", "process_steps"
        } else "labeled_image"
        return Shot(
            template=template_name,
            heading=shot_heading,
            labels=_label_specs(labels, overlay_plan, label_placement),
            arrows=_label_specs(labels, overlay_plan, label_placement),
            highlights=_highlight_specs(overlay_plan),
            motion=motion_data,
        )

    return _shot_from_media_type(media_type, motion, visual, image_prompt, labels)


def _motion_dict(text: str) -> dict[str, str]:
    lowered = (text or "").lower()
    if "arrow_draw_then_label_fade" in lowered or ("arrow" in lowered and "label" in lowered and "fade" in lowered):
        return {"type": "arrow_draw_then_label_fade"}
    if "pan" in lowered:
        return {"type": "pan"}
    if "slide" in lowered:
        return {"type": "slide"}
    if "fade" in lowered:
        return {"type": "fade"}
    if "zoom" in lowered or "ken burns" in lowered:
        return {"type": "zoom"}
    return {"type": "zoom"}


def _plain_items(text: str) -> list[str]:
    return [part.strip(" .:-") for part in re.split(r"[\n;|]+", text or "") if part.strip(" .:-")]


def _formula_lines(text: str) -> list[str]:
    items = _plain_items(text)
    if items:
        return items[:6]
    candidates = re.findall(r"[^.;\n]*(?:=|\\Delta|λ|Λ|C₁|C₂|KCl|NaCl)[^.;\n]*", text or "")
    return [item.strip() for item in candidates if item.strip()][:6]


def _explain_steps(text: str) -> list[str]:
    clauses = [part.strip(" .") for part in re.split(r"\bthen\b|\btherefore\b|\bhence\b|[.;]", text or "", flags=re.I) if part.strip()]
    return [_short_title(part, "Explain") for part in clauses[:4]]


def _columns_from_text(text: str) -> list[dict[str, str]]:
    parts = _plain_items(text)
    if len(parts) >= 2:
        return [{"title": _short_title(parts[0], "Before"), "text": parts[0]}, {"title": _short_title(parts[1], "After"), "text": parts[1]}]
    return [{"title": "Concept", "text": text or "Main visual"}, {"title": "Meaning", "text": "Key learning point"}]


def _label_specs(labels: list[str], overlay_plan: str, label_placement: str = "") -> list[dict[str, str]]:
    names = _expand_label_names(labels or _labels(overlay_plan))
    placement_by_name = _placement_map(label_placement or overlay_plan)
    specs: list[dict[str, str]] = []
    for index, name in enumerate(names):
        label = _short_label(name)
        placement = placement_by_name.get(label.lower()) or placement_by_name.get(name.lower()) or label_placement
        spec = {"text": label, "position": str(index), "placement": placement}
        xy = _coordinate_from_text(placement)
        if xy:
            spec["target_xy"] = xy
        specs.append(spec)
    return specs


def _placement_map(text: str) -> dict[str, str]:
    result: dict[str, str] = {}
    for part in re.split(r"[\n]+", text or ""):
        clean = part.strip(" .")
        if not clean:
            continue
        if "|" in clean:
            name, placement = clean.split("|", 1)
        elif ":" in clean:
            name, placement = clean.split(":", 1)
        else:
            continue
        key = _short_label(name).lower()
        if key and placement.strip():
            result[key] = placement.strip()
    return result


def _coordinate_from_text(text: str) -> str:
    target = re.search(
        r"\btarget\s*=\s*\(\s*(0?\.\d+|1(?:\.0+)?)\s*,\s*(0?\.\d+|1(?:\.0+)?)\s*\)",
        text or "",
        re.I,
    )
    if target:
        return f"{target.group(1)},{target.group(2)}"
    target_xy = re.search(
        r"\btarget_xy\s*:\s*(0?\.\d+|1(?:\.0+)?)\s*,\s*(0?\.\d+|1(?:\.0+)?)",
        text or "",
        re.I,
    )
    if target_xy:
        return f"{target_xy.group(1)},{target_xy.group(2)}"
    match = re.search(r"(?:@|->|\bxy\b|\btarget\b)?\s*(0?\.\d+|1(?:\.0+)?)\s*[,x]\s*(0?\.\d+|1(?:\.0+)?)", text or "", re.I)
    if match:
        return f"{match.group(1)},{match.group(2)}"
    return ""


def _expand_label_names(names: list[str]) -> list[str]:
    result: list[str] = []
    compounds = [
        ("anther stigma", ["anther", "stigma"]),
        ("stigma anther", ["stigma", "anther"]),
        ("anther stigma pollen grains pollen transfer path", ["anther", "stigma", "pollen grains", "pollen transfer path"]),
        ("pollen grains pollen transfer path", ["pollen grains", "pollen transfer path"]),
        ("pollen grains pollinator pollen transfer path", ["pollen grains", "pollinator", "pollen transfer path"]),
        ("wood sorrel small scentless closed", ["wood sorrel", "closed flower"]),
        ("closed wood sorrel flower", ["closed flower", "wood sorrel"]),
        ("bisexual flower self-pollination cleistogamy", ["bisexual flower", "self-pollination", "cleistogamy"]),
        ("nectar guide pollinator", ["nectar guide", "pollinator"]),
    ]
    for name in names:
        normalized = re.sub(r"\s+", " ", (name or "").lower()).strip()
        if _is_empty_overlay_noise(normalized):
            continue
        replacement = next((parts for key, parts in compounds if normalized == key), None)
        if replacement:
            result.extend(replacement)
            continue
        result.append(name)
    seen: set[str] = set()
    unique: list[str] = []
    for item in result:
        key = item.lower().strip()
        if key and key not in seen:
            seen.add(key)
            unique.append(item)
    return unique


def _highlight_specs(overlay_plan: str) -> list[dict[str, str]]:
    return [{"text": _short_label(item), "position": str(index)} for index, item in enumerate(_plain_items(overlay_plan)[:4])]


def _short_label(text: str) -> str:
    text = re.sub(r"(?i)\b(arrows?|highlights?|formula lines?|explain steps?|steps?|columns?|labels?)\b", " ", text or "")
    words = re.findall(r"[A-Za-z0-9][A-Za-z0-9+-]*", text)
    if not words:
        return "Label"
    label = " ".join(words[:3])
    if len(label) <= 22:
        return label
    return words[0][:22]


def _shot_from_media_type(media_type: str, motion: str, visual: str, image_prompt: str, labels: list[str] | None = None) -> Shot | None:
    routing_text = " ".join([media_type, motion, visual, image_prompt]).lower()
    heading = _clean_heading(visual or image_prompt)
    media = media_type.lower()
    is_explicit_process = any(term in " ".join([media_type, motion]).lower() for term in ["animation", "diagram", "local animation"])
    if "wan" in routing_text or "video shot" in routing_text:
        return Shot(template="video", heading=heading or "Video")
    if "photo" in media and not "animation" in media:
        return Shot(template="photo", heading=heading or "Photo")
    if not is_explicit_process and "realistic" in image_prompt.lower():
        return Shot(template="photo", heading=heading or "Photo")
    if "graph" in routing_text:
        return Shot(template="process", heading=heading or "Graph Relationship", steps=["Read the axes", "Compare the trend"], stage_fractions=[0, 0.55])
    if "diagram" in routing_text or "animation" in routing_text or "local animation" in routing_text:
        steps = _steps_from_labels_or_visual(labels or [], visual or image_prompt)
        return Shot(template="process", heading=heading or "Key Process", steps=steps, stage_fractions=[0, 0.55])
    return None


def _clean_heading(text: str) -> str:
    text = re.sub(
        r"\b(animation with labels|local animation|static image|photo with labels|wan video|"
        r"step-by-step animation|quick animation|educational diagram|animated diagram|"
        r"simplified diagram|labeled diagram|diagram showing|animation showing)\b",
        " ",
        text or "",
        flags=re.I,
    )
    text = re.sub(r"\s+", " ", text).strip(" .|-:")
    return _short_title(text, "")


def _display_subject(text: str) -> str:
    text = re.split(r"(?i)\bvisual notes\s*:", text or "", maxsplit=1)[0]
    text = re.split(r"(?i)\blocal animation\s*:", text, maxsplit=1)[0]
    text = re.split(r"(?i)\bcomfy/background prompt\s*:", text, maxsplit=1)[0]
    text = re.sub(r"(?i)\b(real hd educational visual|topic-specific educational visual)\s*:\s*", "", text)
    return re.sub(r"\s+", " ", text).strip(" .:-")


def _steps_from_labels_or_visual(labels: list[str], visual: str) -> list[str]:
    clean_labels = [label.strip() for label in labels if label.strip()]
    if len(clean_labels) >= 2:
        return [_short_title(clean_labels[0], "Observe"), _short_title(clean_labels[1], "Connect")]
    return _two_steps(visual)


def storyboard_from_text(title: str, text: str) -> Storyboard:
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if not paragraphs:
        paragraphs = [text.strip()]
    scenes: list[Scene] = []
    for index, paragraph in enumerate(paragraphs, start=1):
        scene_title = _short_title(paragraph, f"Scene {index}")
        scenes.append(
            Scene(
                scene_number=index,
                title=scene_title,
                narration=paragraph,
                segments=_segments_from_text(paragraph),
            )
        )
    return Storyboard(title=title, source=text, scenes=scenes)


def _segments_from_text(text: str) -> list[Segment]:
    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if s.strip()]
    if not sentences and text.strip():
        sentences = [text.strip()]
    segments: list[Segment] = []
    for index, sentence in enumerate(sentences, start=1):
        keywords = _keywords(sentence)
        visual = "Show a clean educational visual with labeled concepts from the narration."
        if keywords:
            visual = "Show labeled educational visuals for " + ", ".join(keywords[:4]) + "."
        segments.append(
            Segment(
                segment_number=index,
                narration=sentence,
                visual=visual,
                image_prompt=f"Educational illustration about {', '.join(keywords[:5]) or sentence}",
                keywords=keywords,
            )
        )
    return segments


def _keywords(text: str) -> list[str]:
    stop = {
        "about", "above", "after", "again", "also", "because", "before", "being",
        "between", "could", "different", "during", "every", "first", "from",
        "have", "into", "more", "other", "over", "same", "should", "some",
        "such", "than", "that", "their", "them", "then", "there", "these",
        "this", "through", "where", "which", "while", "with", "would", "your",
        "the", "and", "for", "are", "was", "were", "has", "had", "not", "can",
    }
    words = re.findall(r"[A-Za-z][A-Za-z-]{3,}", text.lower())
    seen: set[str] = set()
    result: list[str] = []
    for word in words:
        if word not in stop and word not in seen:
            seen.add(word)
            result.append(word)
    return result[:8]


def _short_title(text: str, fallback: str) -> str:
    words = re.findall(r"[A-Za-z][A-Za-z-]{2,}", text)
    if not words:
        return fallback
    return " ".join(words[:7])


def _storyboard_text(storyboard: Storyboard) -> str:
    return "\n".join(scene.narration for scene in storyboard.scenes if scene.narration)


def _storyboard_text_from_scenes(scenes: list[Scene]) -> str:
    return "\n".join(scene.narration for scene in scenes if scene.narration)


def save_storyboard(storyboard: Storyboard, path: str | Path) -> None:
    Path(path).write_text(
        json.dumps(storyboard.model_dump(), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

