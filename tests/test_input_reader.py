from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from app.input_reader import read_input


def _add_shot_table(document: Document, values: dict[str, str]) -> None:
    table = document.add_table(rows=0, cols=2)
    for key, value in values.items():
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value


class VerticalStoryboardReaderTests(unittest.TestCase):
    def test_reads_vertical_shots_labels_coordinates_and_plain_images(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "lesson_storyboard.docx"
            document = Document()
            document.add_paragraph("Storyboard: General Science Lesson")
            document.add_paragraph("Scene 1: Core concept")
            document.add_paragraph("Narration/Audio/Voiceover:")
            document.add_paragraph("The first concept is labeled. The second is visual context.")

            document.add_paragraph("Shot 1.1: Labeled structure")
            _add_shot_table(document, {
                "Narration": "The first concept is labeled.",
                "Template / Media": "labeled_image\nphoto_with_labels",
                "Tool / Motion Type": "pillow_opencv\nstatic_image",
                "Visual Type / Image Requirement": (
                    "visual_type: realistic_labeled_image\n"
                    "image_requirement: Sharp 1080p subject with visible target parts."
                ),
                "Asset / Image Prompt": "image_prompt: Clean no-text educational image.",
                "Labels": "Part A\nPart B",
                "Label Placement": (
                    "Part A: box=left; target=(0.35,0.42); anchor=visible first part\n"
                    "Part B: box=right; target=(0.62,0.36); anchor=visible second part"
                ),
                "Label Style": "white box, dark outline, yellow arrow",
                "Motion / Subtitle": (
                    "motion: arrow_draw_then_label_fade\n"
                    "subtitle: The first concept is labeled.\n"
                    "subtitle_style: bottom band"
                ),
            })

            document.add_paragraph("Shot 1.2: Context image")
            _add_shot_table(document, {
                "Narration": "The second is visual context.",
                "Template / Media": "labeled_image\nphoto",
                "Tool / Motion Type": "comfy_image\nstatic_image",
                "Visual Type / Image Requirement": (
                    "visual_type: realistic_image\n"
                    "image_requirement: Sharp 1080p context image."
                ),
                "Asset / Image Prompt": "image_prompt: Clean no-text context image.",
                "Labels": "",
                "Label Placement": "",
                "Label Style": "",
                "Motion / Subtitle": "motion: slow_zoom_in\nsubtitle_style: bottom band",
            })
            document.save(path)

            _, storyboard = read_input(path)

            self.assertIsNotNone(storyboard)
            self.assertEqual("General Science Lesson", storyboard.title)
            segments = storyboard.scenes[0].segments
            self.assertEqual(2, len(segments))
            self.assertEqual("realistic_labeled_image", segments[0].shot.template)
            self.assertEqual(["Part A", "Part B"], [item["text"] for item in segments[0].shot.labels])
            self.assertEqual(["0.35,0.42", "0.62,0.36"],
                             [item["target_xy"] for item in segments[0].shot.labels])
            self.assertEqual("bottom band", segments[0].shot.motion["subtitle_style"])
            self.assertEqual("photo", segments[1].shot.template)
            self.assertEqual([], segments[1].shot.labels)


if __name__ == "__main__":
    unittest.main()
