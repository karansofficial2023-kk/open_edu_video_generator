from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from .schema import Storyboard


def export_storyboard_docx(storyboard: Storyboard, output_path: str | Path) -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"Storyboard: {storyboard.title}")
    run.bold = True
    run.font.size = Pt(18)

    for scene in storyboard.scenes:
        header = doc.add_paragraph()
        run = header.add_run(f"Scene {scene.scene_number}: {scene.title}")
        run.bold = True
        run.font.size = Pt(14)

        label = doc.add_paragraph()
        run = label.add_run("Narration/Audio/Voiceover:")
        run.bold = True
        run.italic = True

        narration = doc.add_paragraph()
        narration.paragraph_format.left_indent = Inches(0.25)
        narration.add_run(f'"{scene.narration}"').italic = True

        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        headers = ["S.no", "Splitting the Narration (Sentence wise)", "Visual / Animation", "Image Recommendation"]
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
        for segment in scene.segments:
            row = table.add_row().cells
            row[0].text = str(segment.segment_number)
            row[1].text = f'"{segment.narration}"'
            row[2].text = segment.visual
            row[3].text = segment.image_prompt

    doc.save(output_path)

