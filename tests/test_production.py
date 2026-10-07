from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw

from app import formulas, labels
from app.assets import is_generic_prompt, technical_qa
from app.captions import build_cues
from app.config import AppConfig
from app.contract import PENDING, read_contract_docx, read_contract_json, split_list
from app.preflight import check_storyboard
from app.production_renderer import ShotPlan, ShotRenderer, camera_box
from app.router import clean_title
from app.schema import Caption, Segment
from app.typography import parse_style, subtitle_band

FIELDS = ["shot_id", "narration", "duration", "visual_type", "media_type", "image_requirement", "image_prompt",
          "labels", "label_placement", "label_style", "motion", "subtitle", "asset_path", "review_notes"]


def make_docx(path: Path, shots: list[dict]) -> None:
    doc = Document()
    doc.add_paragraph("Storyboard: Sample Lesson")
    scene = None
    for shot in shots:
        number = shot["shot_id"].split(".")[0]
        if number != scene:
            scene = number
            doc.add_paragraph(f"Scene {number}: Scene title {number}")
            doc.add_paragraph("Narration/Audio/Voiceover:")
            doc.add_paragraph('"Scene narration."')
        doc.add_paragraph(f"Shot {shot['shot_id']}: {shot.get('heading', 'Heading')}")
        rows = [(k, shot[k]) for k in FIELDS if k in shot] + [(k, shot[k]) for k in ("formula_lines", "explain_steps", "steps") if k in shot]
        table = doc.add_table(rows=len(rows), cols=2)
        for row, (key, value) in zip(table.rows, rows):
            row.cells[0].text, row.cells[1].text = key, value
    doc.save(path)


SHOTS = [
    {"shot_id": "1.1", "heading": "Big Topic", "narration": "Welcome to the lesson.", "duration": "4.0 sec", "visual_type": "title_card",
     "image_prompt": "A calm landscape at dawn", "motion": "camera: slow_zoom_in", "subtitle": "Welcome to the lesson."},
    {"shot_id": "1.2", "narration": "The heart pumps blood.", "duration": "6 sec", "visual_type": "realistic_labeled_image",
     "image_prompt": "Anatomical heart photograph", "labels": "Aorta\nLeft ventricle",
     "label_placement": "Aorta - large artery leaving the top\nLeft ventricle - thick muscular lower left chamber", "subtitle": "The heart pumps blood."},
    {"shot_id": "2.1", "narration": "Ohm's law relates voltage and current.", "duration": "5 sec", "visual_type": "formula",
     "formula_lines": "V = I R\nI = V / R", "explain_steps": "Voltage equals current times resistance\nRearrange for current"},
    {"shot_id": "2.2", "narration": "First heat. Then cool.", "duration": "5 sec", "visual_type": "process_steps",
     "steps": "Heat the sample\nCool the sample"},
]


class ContractTests(unittest.TestCase):
    def test_docx_contract_maps_fields_and_routes_renderers(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "lesson.docx"
            make_docx(path, SHOTS)
            board = read_contract_docx(path)
        segments = [seg for _s, seg in board.all_segments()]
        self.assertEqual([s.shot_id for s in segments], ["1.1", "1.2", "2.1", "2.2"])
        self.assertEqual([s.shot.template for s in segments], ["title_card", "photo", "formula", "process"])
        self.assertEqual(segments[0].shot.heading, "Big Topic")
        self.assertEqual(segments[1].labels, ["Aorta", "Left ventricle"])
        self.assertEqual(len(segments[1].label_placement), 2)
        self.assertEqual(segments[2].formula_lines, ["V = I R", "I = V / R"])
        self.assertEqual(segments[3].shot.steps, ["Heat the sample", "Cool the sample"])
        self.assertEqual(segments[0].duration_hint, 4.0)

    def test_json_contract_v2_matches_docx_reader(self):
        data = {"schema_version": "2.0", "title": "T", "language": {"bcp47": "en-IN"},
                "scenes": [{"scene_number": 1, "scene_title": "S", "narration": "N", "shots": [
                    {"shot_id": "1.1", "narration": "Hi.", "duration": 3, "visual_type": "realistic_image", "image_prompt": "p",
                     "labels": [], "label_placement": [], "formula_lines": [], "explain_steps": [], "steps": [], "columns": []}]}]}
        board = read_contract_json(data)
        self.assertEqual(board.language["bcp47"], "en-IN")
        self.assertEqual(board.all_segments()[0][1].shot.template, "photo")

    def test_plain_storyboard_json_is_not_mistaken_for_contract(self):
        self.assertIsNone(read_contract_json({"title": "x", "scenes": [{"scene_number": 1, "segments": [{"narration": "a"}]}]}))

    def test_router_never_uses_topic_keywords(self):
        seg = Segment(segment_number=1, narration="Photosynthesis and pollination and electroplating.", visual_type="realistic_image")
        from app.router import route_segment
        self.assertEqual(route_segment(seg).template, "photo")

    def test_titles_drop_production_metadata(self):
        self.assertNotIn("1920x1080", clean_title("Sharp 1920x1080 production asset Heart anatomy"))

    def test_stock_photo_titles_must_relate_to_the_shot(self):
        from app.assets import title_relevant
        grounding = "This causes the solution to extract silver atoms from the pure silver rod, which acts as the anode. A pure silver rod submerged in an electrolyte solution"
        self.assertFalse(title_relevant("File:Manhès-David vertical converters at the Anaconda Copper.png", grounding))      # real shot 5.5
        self.assertFalse(title_relevant("File:Lemon batteries circuit 1390001 nevit.jpg", "Can you determine which reactions occurs at the anode"))
        self.assertTrue(title_relevant("File:Silver rod in electrolyte.jpg", grounding))
        self.assertTrue(title_relevant("File:Electrolytic Cell Diagram.jpg", "This is an electrolytic cell containing a solution of silver nitrate"))
        self.assertFalse(title_relevant("File:Large concrete floor.jpg", "depict the concrete subject, principal subject large"))     # generic words only

    def test_render_clip_runs_for_text_and_photo_neighbours(self):
        """render_clip was never executed by a test: a wrong attribute path (plan.shot) failed every shot of a real render."""
        import tempfile
        from pathlib import Path
        from PIL import Image
        from app.config import AppConfig
        from app.production_renderer import ShotPlan, ShotRenderer
        from app.schema import Segment, Shot
        config = AppConfig()
        config.output_resolution.width, config.output_resolution.height, config.fps = 960, 540, 12
        with tempfile.TemporaryDirectory() as tmp:
            renderer = ShotRenderer(config, Path(tmp))
            seg = Segment(segment_number=1, narration="Let's turn on the power supply and watch.", start=0.0, end=0.5,
                          shot=Shot(template="process", heading="Power", steps=["Power supply", "Turn it on"], stage_fractions=[0, 0.5]))
            plan = ShotPlan(segment=seg, scene_title="Setup")
            previous = Image.new("RGB", (960, 540), "white")
            for index, (prev_template, expect_fade) in enumerate([(None, False), ("photo", False), ("process", False)]):
                clip = Path(tmp) / f"clip_{index}.mp4"
                last = renderer.render_clip(plan, clip, 6, previous if prev_template else None, prev_template)
                self.assertTrue(clip.exists() and clip.stat().st_size > 500)
                self.assertEqual(last.size, (960, 540))

    def test_reviewer_checklist_must_come_from_the_requirement(self):
        from app.assets import apply_checklist
        required = ("A pure silver rod submerged in an electrolyte solution. This causes the solution to extract silver atoms from the "
                    "pure silver rod, which acts as the anode")
        verdict = {"matches": True, "score": 0.95, "contains_text": False, "missing": [], "issues": []}
        # real shot 5.5: the reviewer described a copper-smelter photograph instead of the required apparatus
        drifted = [{"item": "large industrial apparatus", "visible": True}, {"item": "concrete floor", "visible": True},
                   {"item": "metallic rods or anodes", "visible": True}, {"item": "industrial machinery", "visible": True}]
        result = apply_checklist(dict(verdict), drifted, grounding=required)
        self.assertFalse(result["matches"])
        self.assertLessEqual(result["score"], 0.3)
        # the contract's image prompt is full of boilerplate ("depict the concrete subject ... principal subject large and unobscured"):
        # generic photo words from it must not ground "concrete floor" / "large industrial apparatus"
        boilerplate = required + " Depict the concrete subject or visible action; principal subject large and unobscured, dark background."
        self.assertFalse(apply_checklist(dict(verdict), drifted, grounding=boilerplate)["matches"])
        slide = [{"item": "dark background", "visible": True}, {"item": "white text", "visible": True},
                 {"item": "thin colored border", "visible": True}, {"item": "silver rod", "visible": True}]
        self.assertFalse(apply_checklist(dict(verdict), slide, grounding=boilerplate)["matches"])
        honest = [{"item": "silver rod", "visible": True}, {"item": "solution in a beaker", "visible": True}]
        self.assertTrue(apply_checklist(dict(verdict), honest, grounding=required)["matches"])
        self.assertTrue(apply_checklist(dict(verdict), drifted)["matches"])          # no grounding text: old behaviour

    def test_tts_silence_is_trimmed_and_caption_shift_reported(self):
        import subprocess, tempfile
        from pathlib import Path
        from app import tts_edge
        from app.config import AppConfig
        with tempfile.TemporaryDirectory() as tmp:
            wav = Path(tmp) / "clip.wav"
            # 0.5 s silence + 1 s tone + 1.2 s silence, like an engine that pads both ends
            subprocess.run(["ffmpeg", "-v", "error", "-y", "-f", "lavfi", "-i", "anullsrc=r=24000:cl=mono", "-f", "lavfi", "-i", "sine=f=300:r=24000:d=1",
                            "-f", "lavfi", "-i", "anullsrc=r=24000:cl=mono", "-filter_complex",
                            "[0:a]atrim=0:0.5[a];[1:a]anull[b];[2:a]atrim=0:1.2[c];[a][b][c]concat=n=3:v=0:a=1", "-ar", "24000", "-ac", "1", str(wav)], check=True)
            before = float(subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(wav)],
                                          capture_output=True, text=True).stdout)
            removed = tts_edge.trim_silence(wav, AppConfig())
            after = float(subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(wav)],
                                         capture_output=True, text=True).stdout)
        self.assertAlmostEqual(before, 2.7, delta=0.05)
        self.assertAlmostEqual(removed, 0.5 - tts_edge.KEEP_LEAD, delta=0.06)
        self.assertAlmostEqual(after, 1.0 + tts_edge.KEEP_LEAD + tts_edge.KEEP_TAIL, delta=0.12)     # the tone is never cut

    def test_latex_formula_lines_survive_the_chemistry_rewrites(self):
        """Real contract line from the resistivity lesson: the CH_3COOH rule turned 'A_0}' into 'A0' and the fraction no longer parsed."""
        from app import formulas
        for line in [r"R_0 = \rho \times \frac{L_0}{A_0}", r"A' = \frac{A_0 \times L_0}{L'}", r"R' = 1.21 \times R_0",
                     r"\frac{1}{x^2} = 1 \Rightarrow x = \pm 1"]:
            self.assertIsNone(formulas.validate(line), line)
            formulas.render_line(line, "#1b6b57", 48)
        self.assertIn(r"\frac{L_0}{A_0}", formulas.to_mathtext(r"R_0 = \rho \times \frac{L_0}{A_0}"))
        self.assertIn("CH_{3}", formulas.to_mathtext("CH_3COOH -> CH_3COO- + H+"))         # chemistry still gets its rewrites

    def test_formula_caption_and_note_keep_their_spaces(self):
        from app import formulas
        line = "Redox reaction: Ag⁺ + e⁻ → Ag (reduction at cathode)"
        math = formulas.to_mathtext(line)
        self.assertIn(r"\mathrm{Redox\ reaction:}", math)
        self.assertIn(r"\mathrm{reduction\ at\ cathode}", math)
        self.assertIsNone(formulas.validate(line))
        formulas.render_line(line, "#1b6b57", 48)          # must typeset without raising
        self.assertEqual(formulas.to_mathtext("Ag(s) → Ag⁺(aq) + e⁻"), formulas.to_mathtext("Ag(s) → Ag⁺(aq) + e⁻"))
        self.assertNotIn(r"\ \ (", formulas.to_mathtext("Ag⁺(aq) + e⁻ → Ag(s)"))     # (aq)/(s) are state symbols, not notes

    def test_cut_off_sentence_prefix_is_not_a_heading(self):
        from app.router import route_segment
        narration = "For weak electrolytes, the graph approaches an asymptote at lower concentrations."
        cut = Segment(segment_number=1, narration=narration, heading="For weak electrolytes, the graph approaches an",
                      visual_type="formula", formula_lines=["x = y"])
        self.assertEqual(route_segment(cut).heading, "")
        topic = Segment(segment_number=1, narration=narration, heading="Weak electrolyte graph", visual_type="formula",
                        formula_lines=["x = y"])
        self.assertEqual(route_segment(topic).heading, "Weak electrolyte graph")
        title = Segment(segment_number=1, narration="Applications of Kohlrausch Law and more.", heading="Applications of Kohlrausch Law",
                        visual_type="title_card")
        self.assertEqual(route_segment(title).heading, "Applications of Kohlrausch Law")

    def test_fallback_cards_never_end_on_a_connective(self):
        from app.assets import _split_clauses, _DANGLING
        narrations = [
            "Can you determine whether strong or weak electrolytes exhibit optimal molar conductivity?",
            "Weak acids and bases are classified as weak electrolytes due to their low molar conductivities.",
            "Stay tuned to watch our upcoming videos to learn more.",
            "Now let's explore some more about the variables involved in this process.",
            "Let's turn on the power supply.",
            "Keep learning, we believe in you.",
            "Do you think there are other applications for Kohlrausch's law of independent migration of ions?",
            "Let's begin by reading the question carefully: Two sources of equal electromotive force (EMF) are connected to an external resistance R.",
        ]
        for text in narrations:
            cards = _split_clauses(text)
            self.assertEqual(len(cards), 2, cards)
            if len(text.split()) < 8:            # short sentence: a key-term card made of its own words, then the whole sentence
                self.assertEqual(cards[1].lower(), text.rstrip(".?!").lower())
                self.assertIn(cards[0].lower(), text.lower())
                self.assertLessEqual(len(cards[0].split()), 3)
                continue
            first, second = cards
            self.assertNotIn(first.split()[-1].lower().strip(",;:"), _DANGLING, (first, second))
            self.assertNotIn(second.split()[0].lower(), {"or", "of", "as", "to", "than"}, (first, second))
            self.assertEqual(" ".join((first + " " + second).split()).lower().replace(",", "").replace(":", ""),
                             text.rstrip(".?!").lower().replace(",", "").replace(":", ""))   # nothing lost, nothing added

    def test_cards_with_invented_facts_are_not_faithful(self):
        from app.assets import _faithful
        narration = "Can you determine whether strong or weak electrolytes exhibit optimal molar conductivity?"
        self.assertTrue(_faithful("Electrolytes: weak electrolytes exhibit molar conductivity", narration))
        self.assertFalse(_faithful("Strong electrolytes: they dissociate completely into free ions", narration))

    def test_split_list(self):
        self.assertEqual(split_list("- a\n2. b;c"), ["a", "b", "c"])


class FormulaTests(unittest.TestCase):
    def test_valid_and_invalid_formulas(self):
        for good in ["V = I R", "CuSO4 + Fe -> FeSO4 + Cu", "Cu^{2+} + 2e^- -> Cu(s)", "a^2 + b^2 = c^2"]:
            self.assertIsNone(formulas.validate(good), good)
        self.assertIn("dangling", formulas.validate("lambda^0_"))
        self.assertIn("prose", formulas.validate("Electrons flow from the anode to the cathode in the external circuit"))
        self.assertIn("brackets", formulas.validate("(a + b"))

    def test_render_returns_transparent_image(self):
        image = formulas.render_line("E = m c^2", "#123456", 60)
        self.assertEqual(image.mode, "RGBA")
        self.assertGreater(image.width, 100)
        low, high = image.getchannel("A").getextrema()
        self.assertLess(low, 20)            # background transparent, not an opaque block
        self.assertGreater(high, 200)       # glyphs opaque


class SubtitleTests(unittest.TestCase):
    def test_band_fits_two_lines_inside_safe_margins(self):
        config = AppConfig()
        text = "While bees are the primary pollinators of sunflowers, these plants have a clever mechanism to ensure pollination when bees are absent."
        overlay, band = subtitle_band((1920, 1080), text, config.production.subtitle)
        left_margin = round(1920 * config.production.subtitle.margin_x)
        self.assertGreaterEqual(band[0], left_margin - 20)
        self.assertLessEqual(band[2], 1920 - left_margin + 20)
        self.assertLessEqual(band[3], 1080 - round(1080 * config.production.subtitle.margin_bottom))
        self.assertLess(band[3] - band[1], 200)   # two lines, not an oversized band

    def test_cues_all_fit_and_cover_speech(self):
        config = AppConfig()
        seg = Segment(segment_number=1, narration=" ".join(["word"] * 60), start=0, end=12, speech_duration=12)
        cues = build_cues(seg, config.production.subtitle, (1920, 1080))
        self.assertGreater(len(cues), 1)
        for cue in cues:
            subtitle_band((1920, 1080), cue.text, config.production.subtitle)
        self.assertEqual(" ".join(c.text for c in cues), seg.narration)

    def test_impossible_line_raises_instead_of_clipping(self):
        with self.assertRaises(ValueError):
            subtitle_band((640, 360), "x" * 400, AppConfig().production.subtitle.model_copy(update={"font_px": 200}))

    def test_style_parser(self):
        self.assertEqual(parse_style("band_color=black; max_lines=2; bold")["max_lines"], "2")


class LabelTests(unittest.TestCase):
    def test_targets_strip_label_prefix(self):
        seg = Segment(segment_number=1, narration="x", labels=["Aorta", "Valve"],
                      label_placement=["Aorta - large artery", PENDING])
        self.assertEqual(labels.label_targets(seg), [("Aorta", "large artery"), ("Valve", "Valve")])

    def test_layout_is_collision_free_and_in_bounds(self):
        config = AppConfig()
        anchors = [labels.Anchor("Left", "d", (0.2, 0.4, 0.28, 0.5), 0.9), labels.Anchor("Right", "d", (0.7, 0.4, 0.78, 0.5), 0.9),
                   labels.Anchor("Middle", "d", (0.45, 0.3, 0.55, 0.4), 0.9)]
        keepout = [(0, 864, 1920, 1080)]
        placed = labels.layout(anchors, (1920, 1080), keepout, config, {})
        self.assertEqual(len(placed), 3)
        for i, a in enumerate(placed):
            self.assertGreaterEqual(a.box[0], 0)
            self.assertLessEqual(a.box[2], 1920)
            self.assertFalse(labels._overlap(a.box, keepout[0]))
            for b in placed[i + 1:]:
                self.assertFalse(labels._overlap(a.box, b.box))
            self.assertGreaterEqual(a.box[3] - a.box[1], 28)   # 28 px minimum text at 1080p

    def test_arrow_precedes_label_text(self):
        config = AppConfig()
        placed = labels.layout([labels.Anchor("Part", "d", (0.4, 0.4, 0.5, 0.5), 0.9)], (1920, 1080), [], config, {})
        early = Image.new("RGBA", (1920, 1080))
        labels.draw(early, placed, 0.2, [0.0], config, {})
        self.assertIsNotNone(early.getbbox())              # line already growing
        box = placed[0].box
        self.assertIsNone(early.crop(box).point(lambda v: 0).getbbox() and None)
        text_region = early.crop((box[0] + 4, box[1] + 4, box[2] - 4, box[3] - 4)).getchannel("A").getbbox()
        self.assertIsNone(text_region)                     # label box not yet visible at t=0.2
        late = Image.new("RGBA", (1920, 1080))
        labels.draw(late, placed, 1.5, [0.0], config, {})
        self.assertIsNotNone(late.crop(box).getchannel("A").getbbox())

    def test_low_confidence_or_invalid_boxes_are_rejected(self):
        self.assertIsNone(labels._norm_box([500, 500, 400, 600]))
        self.assertIsNone(labels._norm_box([0, 0, 1000, 1000]))
        self.assertEqual(labels._norm_box([100, 200, 300, 400]), (0.1, 0.2, 0.3, 0.4))


class AssetTests(unittest.TestCase):
    def test_generic_prompt_detection(self):
        seg = Segment(segment_number=1, narration="The mitochondria produce energy for the cell.",
                      image_prompt="Sharp 1920x1080 production asset in the subject-appropriate visual medium")
        self.assertTrue(is_generic_prompt(seg))
        seg.image_prompt = "Close photograph of a mitochondria inside a cell producing energy"
        self.assertFalse(is_generic_prompt(seg))

    def test_technical_qa_flags_blank_and_blurry(self):
        config = AppConfig()
        self.assertIn("blank_or_flat", technical_qa(Image.new("RGB", (1920, 1080), "white"), config)["reasons"])
        noisy = Image.effect_noise((1920, 1080), 60).convert("RGB")
        self.assertTrue(technical_qa(noisy, config)["ok"])


class PreflightTests(unittest.TestCase):
    def _board(self, **overrides):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "l.docx"
            shots = [dict(SHOTS[1], **overrides)]
            make_docx(path, shots)
            return read_contract_docx(path)

    def test_label_count_mismatch_and_motion_labels_are_errors(self):
        board = self._board(label_placement="Aorta - artery")
        fields = {(i["field"]) for i in check_storyboard(board, AppConfig()) if i["severity"] == "error"}
        self.assertIn("label_placement", fields)
        board = self._board(visual_type="short_motion_clip")
        board.all_segments()[0][1].motion_prompt = "clip"
        board.all_segments()[0][1].shot.template = "video"
        reasons = " ".join(i["reason"] for i in check_storyboard(board, AppConfig()))
        self.assertIn("cannot be anchored", reasons)

    def test_prose_formula_is_reported_with_repair(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "l.docx"
            make_docx(path, [dict(SHOTS[2], formula_lines="Voltage is proportional to the current through the resistor")])
            board = read_contract_docx(path)
        issues = [i for i in check_storyboard(board, AppConfig()) if i["field"] == "formula_lines"]
        self.assertTrue(issues and "explain_steps" in issues[0]["repair"])
        self.assertEqual(set(issues[0]), {"shot_id", "field", "severity", "renderer", "reason", "repair"})

    def test_metadata_in_subtitle_is_blocked(self):
        board = self._board(subtitle="Sharp 1920x1080 production asset")
        self.assertTrue(any(i["field"] == "subtitle" and i["severity"] == "error" for i in check_storyboard(board, AppConfig())))


class RendererTests(unittest.TestCase):
    def test_ken_burns_is_static_when_labels_need_fixed_anchors(self):
        self.assertEqual(camera_box("camera: slow_zoom_in", 3, 6, (1920, 1080), False), (0.0, 0.0, 1920.0, 1080.0))
        box = camera_box("camera: slow_zoom_in", 6, 6, (1920, 1080), True)
        self.assertLess(box[2] - box[0], 1920)

    def test_formula_and_title_frames_render_without_overflow(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "l.docx"
            make_docx(path, SHOTS)
            board = read_contract_docx(path)
            config = AppConfig()
            config.output_resolution.width, config.output_resolution.height = 1920, 1080
            renderer = ShotRenderer(config, Path(folder))
            seg = board.all_segments()[2][1]
            seg.start, seg.end = 0, 9
            frame = renderer.base_frame(ShotPlan(segment=seg, scene_title="Ohm"), 6.0, 9.0)
            self.assertEqual(frame.size, (1920, 1080))
            title = board.all_segments()[0][1]
            image = Path(folder) / "bg.png"
            Image.new("RGB", (1920, 1080), "#557799").save(image)
            title.shot.asset_path = str(image)
            plan = ShotPlan(segment=title, scene_title="Scene")
            renderer.prepare(plan)
            frame = renderer.base_frame(plan, 2.0, 4.0)
            self.assertNotEqual(frame.getpixel((960, 540)), (85, 119, 153))   # title text drawn over the background


class RoutingAndAssetTests(unittest.TestCase):
    def test_fake_steps_cut_from_narration_are_not_a_process(self):
        from app.router import route_segment
        narration = "Fossils serve as one of the most compelling forms of biological evidence, preserving the remains of ancient organisms."
        seg = Segment(segment_number=1, narration=narration, visual_type="process_steps",
                      steps=["Fossils serve as one of the most compelling forms of biological evidence", "preserving the remains of ancient organisms."])
        self.assertEqual(route_segment(seg).template, "photo")
        seg.steps = ["Collect sample", "Heat sample", "Record result"]
        self.assertEqual(route_segment(seg).template, "process")

    def test_prompt_cleaner_removes_templates_and_text_inducing_words(self):
        from app.assets import clean_prompt
        cleaned = clean_prompt("Premium educational documentary frame, sharp native 1920x1080 detail. A title card featuring a fossil with labels. No generated text.")
        self.assertNotIn("1920", cleaned)
        self.assertNotIn("title", cleaned.lower())
        self.assertNotIn("labels", cleaned.lower())
        self.assertIn("fossil", cleaned)

    def test_split_screen_creates_one_job_per_panel(self):
        from app.assets import StillProducer
        from app.schema import Scene, Shot, Storyboard
        seg = Segment(segment_number=1, narration="Compare A and B.", shot_id="1.1", visual_type="split_screen",
                      columns=["Wind | grass releasing pollen in a breeze", "Water | aquatic plant with floating pollen"],
                      shot=Shot(template="split_screen", steps=["Wind", "Water"]))
        board = Storyboard(title="T", scenes=[Scene(scene_number=1, title="S", narration="N", segments=[seg])])
        with tempfile.TemporaryDirectory() as folder:
            jobs = StillProducer(board, Path(folder), AppConfig())._jobs()
        self.assertEqual([j.key for j in jobs], ["1_1_p1", "1_1_p2"])
        self.assertIn("grass", jobs[0].segment.image_requirement)

    def test_graph_rows_and_preflight(self):
        from app.charts import parse_rows
        self.assertEqual(parse_rows(["A: 10%", "B = 2,500", "not a row"]), [("A", 10.0, "%"), ("B", 2500.0, "")])
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "l.docx"
            make_docx(path, [{"shot_id": "1.1", "narration": "Data.", "visual_type": "graph", "duration": "6 sec"}])
            board = read_contract_docx(path)
        reasons = " ".join(i["reason"] for i in check_storyboard(board, AppConfig()))
        self.assertIn("at least two", reasons)
        self.assertIn("source provenance", reasons)


if __name__ == "__main__":
    unittest.main()

